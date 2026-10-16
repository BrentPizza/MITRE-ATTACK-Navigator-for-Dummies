"""
MITRE ATT&CK Navigator for Dummies
----------------------------------
Pick a set of APT groups -> get the ATT&CK techniques (T-codes) they use,
color-coded by how many of the selected groups overlap on each one. The color scale
grows with the number of groups selected (see color_bands()):

    groups  LIGHT GREEN  YELLOW        ORANGE         RED              DARK RED
    1       1            -             -              -                -
    2       1            -             -              2                -
    3       1            2             -              3                -
    4       1            2             3              4                -
    5+      1            the counts from 2 to (groups - 1) split     ALL selected
                         into three near-equal runs
            e.g. 6 groups:  green 1, yellow 2-3, orange 4, red 5, dark red 6
                 10 groups: green 1, yellow 2-4, orange 5-7, red 8-9, dark red 10

Outputs (per run, in ./output/<timestamp>_<matrix>_<groups>/):
    *_legend_for-light-navigator.svg / *_legend_for-dark-navigator.svg -> slide legends (+ .png if Pillow is installed)
and two output sets:
    1_Parent_Only/     parent T-codes only (sub-technique use rolled up into the parent)
    2_Parent_and_Sub/  parent T-codes plus the sub-techniques used, each with its own color
each containing:
    *_layer.json   -> import into MITRE ATT&CK Navigator ("Open Existing Layer")
    *.xlsx         -> Excel workbook (summary, overlap table, tactic matrix)
    *.pdf          -> dark-themed PDF report with the matrix view

Usage:  python attack_overlap.py [mobile|ics] <groups...> [options]
    Enterprise is the default matrix. Put "mobile" or "ics" first to use that matrix instead.

    python attack_overlap.py                                   (interactive, dummy-proof mode)
    python attack_overlap.py APT29 "Aquatic Panda" APT1        (Enterprise)
    python attack_overlap.py mobile APT28 APT-C-23 Windshift
    python attack_overlap.py ics "Sandworm Team" ALLANITE APT33
    python attack_overlap.py G0007 G0016 --no-campaigns
    python attack_overlap.py --list                            (list every Enterprise group)
    python attack_overlap.py mobile --list                     (list every Mobile group)
    python attack_overlap.py ics --search lazarus              (search names/aliases)
    python attack_overlap.py --update APT28 APT29              (re-download latest ATT&CK data)
    python attack_overlap.py --offline APT28 APT29             (no network: use the data bundled in data/)
    python attack_overlap.py APT28 APT29 --mode parents        (only the parent-only set)
"""

import argparse
import datetime as dt
import difflib
import json
import os
import re
import sys
import urllib.request
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
DATA_DIR = HERE / "data"

# openpyxl and reportlab ship in vendor/ as plain pure-Python source, so the tool runs on a stock Python
# with nothing installed. Appended (not prepended) so copies already installed on the machine take priority.
if (HERE / "vendor").is_dir():
    sys.path.append(str(HERE / "vendor"))

# Pillow is optional (it is compiled code, which locked-down workstations may block). Without it every output
# is still produced; the slide legend is written as SVG only, and reportlab gets a stand-in for the PIL module
# it imports at load time but only uses to embed pictures (the PDFs here have none).
try:
    import PIL  # noqa: F401
    HAVE_PILLOW = True
except ImportError:
    HAVE_PILLOW = False
    sys.path.append(str(HERE / "vendor" / "_no_pillow"))
OUT_DIR = HERE / "output"

DOMAINS = {  # matrix -> (Navigator domain / data file name, download URL)
    "enterprise": ("enterprise-attack", "https://raw.githubusercontent.com/mitre/cti/master/enterprise-attack/enterprise-attack.json"),
    "mobile": ("mobile-attack", "https://raw.githubusercontent.com/mitre/cti/master/mobile-attack/mobile-attack.json"),
    "ics": ("ics-attack", "https://raw.githubusercontent.com/mitre/cti/master/ics-attack/ics-attack.json"),
}
DOMAIN_LABEL = {"enterprise": "Enterprise", "mobile": "Mobile", "ics": "ICS"}
DOMAIN_ALIASES = {"enterprise": "enterprise", "enterprise-attack": "enterprise", "ent": "enterprise",
                  "mobile": "mobile", "mobile-attack": "mobile", "mob": "mobile",
                  "ics": "ics", "ics-attack": "ics"}


def parse_domain(text):
    return DOMAIN_ALIASES.get(text.strip().lower())

# Overlap levels, strongest first: key -> (short name, hex color, text color on that fill)
ALL, RED, ORANGE, YELLOW, GREEN = "all", "high", "medium", "low", "single"
LEVELS = {
    ALL: ("DARK RED", "#9D0208", "#FFFFFF"),
    RED: ("RED", "#E5383B", "#FFFFFF"),
    ORANGE: ("ORANGE", "#F48C06", "#1A1A1A"),
    YELLOW: ("YELLOW", "#F4D35E", "#1A1A1A"),
    GREEN: ("LIGHT GREEN", "#90EE90", "#1A1A1A"),
}
LEVEL_SHORT = {k: v[0] for k, v in LEVELS.items()}
LEVEL_ORDER = []   # levels in use for this run, strongest first; filled by set_bands()
BANDS = []         # (level, min groups, max groups) for this run; filled by set_bands()


# --------------------------------------------------------------------------- data

def load_bundle(domain="enterprise", update=False, offline=False):
    """Load the ATT&CK STIX bundle from data/ (the repo ships all three matrices).

    offline: never touch the network; fail with a clear message if the file is missing.
    update:  re-download the latest bundle from mitre/cti first.
    """
    name, url = DOMAINS[domain]
    path = DATA_DIR / f"{name}.json"
    if offline:
        if not path.exists():
            sys.exit(f"Offline mode: {path} is missing. Copy it from the repo (data/{name}.json) "
                     f"or run once with internet access.")
    elif update or not path.exists():
        DATA_DIR.mkdir(exist_ok=True)
        print(f"Downloading latest MITRE ATT&CK {DOMAIN_LABEL[domain]} data...")
        tmp = path.with_suffix(".tmp")
        urllib.request.urlretrieve(url, tmp)
        tmp.replace(path)
    with open(path, encoding="utf-8") as f:
        return json.load(f), name, path


def ext_id(obj):
    for ref in obj.get("external_references", []):
        if ref.get("source_name") in ("mitre-attack", "mitre-mobile-attack", "mitre-ics-attack") and ref.get("external_id"):
            return ref["external_id"], ref.get("url", "")
    return None, ""


def live(obj):
    return not obj.get("revoked") and not obj.get("x_mitre_deprecated")


class Attack:
    def __init__(self, bundle, domain_name):
        self.domain_name = domain_name
        objs = [o for o in bundle["objects"] if live(o)]
        by_id = {o["id"]: o for o in objs}

        matrix = next(o for o in objs if o["type"] == "x-mitre-matrix")
        self.data_modified = matrix.get("modified", "")[:10]
        tactics = {o["id"]: o for o in objs if o["type"] == "x-mitre-tactic"}
        self.tactics = [(tactics[r]["x_mitre_shortname"], tactics[r]["name"]) for r in matrix["tactic_refs"] if r in tactics]
        self.tactic_name = dict(self.tactics)

        # techniques
        self.tech = {}        # stix id -> dict
        self.parents = {}     # T-code -> dict (parent techniques only)
        for o in objs:
            if o["type"] != "attack-pattern":
                continue
            tid, url = ext_id(o)
            if not tid:
                continue
            t = {
                "id": tid,
                "name": o["name"],
                "url": url,
                "tactics": [p["phase_name"] for p in o.get("kill_chain_phases", []) if p.get("kill_chain_name", "").startswith("mitre")],
                "is_sub": bool(o.get("x_mitre_is_subtechnique")) or "." in tid,
                "platforms": o.get("x_mitre_platforms", []),
            }
            self.tech[o["id"]] = t
            if not t["is_sub"]:
                self.parents[tid] = t
        self.by_tid = {t["id"]: t for t in self.tech.values()}
        self.sub_ids = sorted(t["id"] for t in self.tech.values() if t["is_sub"])
        self.platforms = sorted({p for t in self.parents.values() for p in t["platforms"]})

        # groups
        self.groups = {}      # stix id -> dict
        for o in objs:
            if o["type"] != "intrusion-set":
                continue
            gid, url = ext_id(o)
            if gid:
                self.groups[o["id"]] = {"stix": o["id"], "id": gid, "name": o["name"], "url": url,
                                        "aliases": [a for a in o.get("aliases", []) if a != o["name"]]}

        # relationships
        self.uses = defaultdict(set)          # source stix -> set(target stix)
        self.campaigns_of = defaultdict(set)  # group stix -> set(campaign stix)
        for o in objs:
            if o["type"] != "relationship" or o["source_ref"] not in by_id or o["target_ref"] not in by_id:
                continue
            if o["relationship_type"] == "uses":
                self.uses[o["source_ref"]].add(o["target_ref"])
            elif o["relationship_type"] == "attributed-to" and o["target_ref"] in self.groups:
                self.campaigns_of[o["target_ref"]].add(o["source_ref"])
        self.software = {o["id"] for o in objs if o["type"] in ("malware", "tool")}

    # -- lookup
    def find_group(self, query):
        """Return (group, None) or (None, [suggestions])."""
        q = query.strip().lower()
        squash = lambda s: re.sub(r"[^a-z0-9]", "", s.lower())  # "AquaticPanda" == "Aquatic Panda"
        sq = squash(q)
        exact = [g for g in self.groups.values()
                 if sq and sq in {squash(x) for x in [g["id"], g["name"]] + g["aliases"]}]
        if len(exact) == 1:
            return exact[0], None
        if len(exact) > 1:
            return None, exact
        partial = [g for g in self.groups.values()
                   if q in g["name"].lower() or any(q in a.lower() for a in g["aliases"])]
        if len(partial) == 1:
            return partial[0], None
        if partial:
            return None, partial
        names = {g["name"].lower(): g for g in self.groups.values()}
        for g in self.groups.values():
            for a in g["aliases"]:
                names.setdefault(a.lower(), g)
        close = difflib.get_close_matches(q, names.keys(), n=6, cutoff=0.6)
        return None, list({names[c]["id"]: names[c] for c in close}.values())

    def search(self, text):
        q = text.lower()
        return [g for g in sorted(self.groups.values(), key=lambda g: g["name"].lower())
                if q in g["name"].lower() or q in g["id"].lower() or any(q in a.lower() for a in g["aliases"])]

    # -- technique usage for one group, rolled up to parents
    def group_techniques(self, group, campaigns=True, software=False):
        """Return {parent T-code: set(sub-technique IDs observed)}."""
        sources = {group["stix"]}
        if campaigns:
            sources |= self.campaigns_of[group["stix"]]
        targets = set()
        for s in list(sources):
            targets |= self.uses[s]
            if software:
                for sw in self.uses[s] & self.software:
                    targets |= self.uses[sw]
        out = defaultdict(set)
        for t in targets:
            tech = self.tech.get(t)
            if not tech:
                continue
            parent = tech["id"].split(".")[0]
            if parent not in self.parents:
                continue
            out[parent]  # make sure parent key exists even with no subs
            if tech["is_sub"]:
                out[parent].add(tech["id"])
        return out


# --------------------------------------------------------------------------- analysis

def color_bands(n):
    """Color scale for n selected groups: [(level, min groups, max groups)], strongest first.

    1 group is always LIGHT GREEN. Up to 4 groups every count gets its own color (green, yellow,
    orange, red). From 5 groups on, the counts from 2 to n-1 are split into three near-equal runs
    (yellow, orange, red) and DARK RED is kept for techniques used by every selected group.
    """
    if n <= 1:
        return [(GREEN, 1, 1)]
    if n <= 4:
        ladder = {2: [GREEN, RED], 3: [GREEN, YELLOW, RED], 4: [GREEN, YELLOW, ORANGE, RED]}[n]
        return [(lv, c, c) for c, lv in reversed(list(enumerate(ladder, 1)))]
    # Split the counts between "one group" and "all groups" (2 .. n-1) into three near-equal runs.
    mid = n - 2
    sizes = [mid // 3 + (1 if i < mid % 3 else 0) for i in range(3)]   # yellow, orange, red
    y_hi = 1 + sizes[0]
    o_hi = y_hi + sizes[1]
    bands = [(ALL, n, n), (RED, o_hi + 1, n - 1), (ORANGE, y_hi + 1, o_hi), (YELLOW, 2, y_hi), (GREEN, 1, 1)]
    return [b for b in bands if b[1] <= b[2]]


def set_bands(n):
    BANDS[:] = color_bands(n)
    LEVEL_ORDER[:] = [lv for lv, _, _ in BANDS]


def overlap_level(count, n):
    for lv, lo, hi in BANDS:
        if lo <= count <= hi:
            return lv
    raise ValueError(f"no color band for {count}/{n}")


def band_text(lo, hi, n):
    if n == 1:
        return "Used by the selected group"
    if lo == hi == 1:
        return "Used by only ONE group"
    if lo == hi == n:
        return "Used by BOTH groups" if n == 2 else f"Used by ALL {n} groups"
    if lo == hi:
        return f"Used by exactly {lo} of {n} groups"
    return f"Used by {lo} to {hi} of {n} groups"


def level_meanings(n):
    return {lv: band_text(lo, hi, n) for lv, lo, hi in BANDS}


def tid_key(tid):
    p, _, s = tid[1:].partition(".")
    return int(p), int(s or 0)


def analyze(attack, groups, campaigns=True, software=False):
    """Rows = used parent T-codes; each row carries "sub_rows" = used sub-techniques with their own overlap level."""
    set_bands(len(groups))
    per_group = {g["id"]: attack.group_techniques(g, campaigns, software) for g in groups}
    n = len(groups)
    rows = []
    all_tids = sorted({t for usage in per_group.values() for t in usage}, key=tid_key)
    for tid in all_tids:
        users = [g for g in groups if tid in per_group[g["id"]]]
        subs = sorted(set().union(*(per_group[g["id"]][tid] for g in users)), key=tid_key)
        sub_rows = []
        for sid in subs:
            s_users = [g["id"] for g in users if sid in per_group[g["id"]][tid]]
            st = attack.by_tid[sid]
            sub_rows.append({"id": sid, "name": st["name"], "url": st["url"], "count": len(s_users),
                             "level": overlap_level(len(s_users), n), "groups": s_users})
        sub_rows.sort(key=lambda s: (-s["count"], tid_key(s["id"])))
        t = attack.parents[tid]
        rows.append({
            "id": tid, "name": t["name"], "url": t["url"],
            "tactics": [tac for tac, _ in attack.tactics if tac in t["tactics"]],
            "count": len(users), "level": overlap_level(len(users), n),
            "groups": [g["id"] for g in users],
            "subs": subs,
            "sub_rows": sub_rows,
        })
    rows.sort(key=lambda r: (-r["count"], tid_key(r["id"])))
    per_group_counts = {gid: len(u) for gid, u in per_group.items()}
    return rows, per_group_counts


def level_counts(rows):
    """{level: (parent count, sub-technique count)}"""
    return {k: (sum(1 for r in rows if r["level"] == k),
                sum(1 for r in rows for s in r["sub_rows"] if s["level"] == k)) for k in LEVEL_ORDER}


# --------------------------------------------------------------------------- outputs

def safe_name(s):
    return re.sub(r"[^A-Za-z0-9_-]+", "-", s).strip("-")


def write_layer(path, attack, groups, rows, title, opts):
    n = len(groups)
    gname = {g["id"]: g["name"] for g in groups}
    used = {r["id"] for r in rows}
    used_subs = {s["id"] for r in rows for s in r["sub_rows"]}

    def entry(t, **extra):
        return {
            "techniqueID": t["id"],
            "score": t["count"],
            "color": LEVELS[t["level"]][1],
            "comment": f"{LEVEL_SHORT[t['level']]}: used by {t['count']}/{n} - " + ", ".join(gname[g] for g in t["groups"]),
            "enabled": True,
            "metadata": [{"name": "Groups", "value": ", ".join(gname[g] for g in t["groups"])}],
            **extra,
        }

    techniques = []
    for r in rows:
        techniques.append(entry(r, showSubtechniques=bool(r["sub_rows"])))
        techniques += [entry(s, showSubtechniques=False) for s in r["sub_rows"]]
    # Disable every unused parent AND unused sub-technique so "hide disabled" shows only what the groups use.
    # (Navigator keeps a parent visible while any of its sub-techniques is still enabled.)
    for tid in sorted(attack.parents, key=tid_key):
        if tid not in used:
            techniques.append({"techniqueID": tid, "enabled": False, "showSubtechniques": False})
    for sid in attack.sub_ids:
        if sid not in used_subs:
            techniques.append({"techniqueID": sid, "enabled": False, "showSubtechniques": False})

    layer = {
        "name": title[:100],
        "versions": {"layer": "4.5", "navigator": "5.1.0"},
        "domain": attack.domain_name,
        "description": (f"{'Technique and sub-technique' if opts['subs'] else 'Parent technique'} overlap for: "
                        f"{', '.join(g['name'] + ' (' + g['id'] + ')' for g in groups)}. "
                        + "; ".join(f"{LEVEL_SHORT[k]} = {m.replace('Used by ', '')}" for k, m in level_meanings(n).items()) + ". "
                        + ("Parents are colored by rolled-up use; sub-techniques by their own use. " if opts["subs"]
                           else "Sub-technique use is rolled up into the parent T-code. ") +
                        f"Campaigns {'included' if opts['campaigns'] else 'excluded'}; "
                        f"software-derived techniques {'included' if opts['software'] else 'excluded'}. "
                        f"ATT&CK data modified {attack.data_modified}."),
        "filters": {"platforms": attack.platforms},
        "sorting": 3,
        "layout": {"layout": "side", "aggregateFunction": "average", "showID": True, "showName": True,
                   "showAggregateScores": False, "countUnscored": False,
                   "expandedSubtechniques": "annotated" if opts["subs"] else "none"},
        "hideDisabled": True,
        "techniques": techniques,
        "gradient": {"colors": [LEVELS[k][1] for k in reversed(LEVEL_ORDER)] if len(LEVEL_ORDER) > 1
                     else [LEVELS[GREEN][1]] * 2, "minValue": 1, "maxValue": max(n, 2)},
        "legendItems": [{"label": f"{LEVEL_SHORT[k]} - {m}", "color": LEVELS[k][1]} for k, m in level_meanings(n).items()],
        "showTacticRowBackground": True,
        "tacticRowBackground": "#1F2633",
        "selectTechniquesAcrossTactics": True,
        "selectSubtechniquesWithParent": False,
        "selectVisibleTechniques": False,
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(layer, f, indent=2)


def write_excel(path, attack, groups, rows, per_group_counts, title, opts):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def fill(hex_):
        return PatternFill("solid", start_color=hex_.lstrip("#"), end_color=hex_.lstrip("#"))

    HDR = fill("#1F2633")
    HFONT = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="3A4250")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    WRAP = Alignment(wrap_text=True, vertical="top")
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

    n = len(groups)
    counts = level_counts(rows)
    n_subs = sum(len(r["sub_rows"]) for r in rows)
    wb = Workbook()

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Generated {dt.datetime.now():%Y-%m-%d %H:%M} | ATT&CK {attack.domain_name} data modified {attack.data_modified} | " \
               f"Campaigns {'included' if opts['campaigns'] else 'excluded'} | Software-derived {'included' if opts['software'] else 'excluded'}"
    ws["A2"].font = Font(italic=True, color="666666")

    ws["A4"] = "Selected Groups"
    ws["A4"].font = Font(bold=True, size=12)
    hdr = ["Group ID", "Name", "Aliases", "Parent T-codes used", "ATT&CK URL"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.fill, cell.font, cell.border = HDR, HFONT, BORDER
    for i, g in enumerate(groups, 6):
        vals = [g["id"], g["name"], ", ".join(g["aliases"]), per_group_counts[g["id"]], g["url"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border, cell.alignment = BORDER, WRAP
        ws.cell(row=i, column=5).hyperlink = g["url"]

    r0 = 6 + n + 1
    ws.cell(row=r0, column=1, value="Color Legend").font = Font(bold=True, size=12)
    legend_cols = 4 if opts["subs"] else 3
    for c, h in enumerate(["Color", "Meaning", "Technique count", "Sub-technique count"][:legend_cols], 1):
        cell = ws.cell(row=r0 + 1, column=c, value=h)
        cell.fill, cell.font, cell.border = HDR, HFONT, BORDER
    meaning = level_meanings(n)
    for i, k in enumerate(LEVEL_ORDER, r0 + 2):
        a = ws.cell(row=i, column=1, value=LEVEL_SHORT[k])
        a.fill, a.font = fill(LEVELS[k][1]), Font(bold=True, color=LEVELS[k][2].lstrip("#"))
        ws.cell(row=i, column=2, value=meaning[k])
        ws.cell(row=i, column=3, value=counts[k][0])
        if opts["subs"]:
            ws.cell(row=i, column=4, value=counts[k][1])
        for c in range(1, legend_cols + 1):
            ws.cell(row=i, column=c).border = BORDER
    tot = r0 + 2 + len(LEVEL_ORDER)
    ws.cell(row=tot, column=2, value="Total unique T-codes").font = Font(bold=True)
    ws.cell(row=tot, column=3, value=len(rows)).font = Font(bold=True)
    if opts["subs"]:
        ws.cell(row=tot, column=4, value=n_subs).font = Font(bold=True)

    t0 = tot + 2
    ws.cell(row=t0, column=1, value="By Tactic").font = Font(bold=True, size=12)
    for c, h in enumerate(["Tactic"] + [LEVEL_SHORT[k] for k in LEVEL_ORDER] + ["Total"], 1):
        cell = ws.cell(row=t0 + 1, column=c, value=h)
        cell.fill, cell.font, cell.border = HDR, HFONT, BORDER
    for k, col in zip(LEVEL_ORDER, range(2, 2 + len(LEVEL_ORDER))):
        ws.cell(row=t0 + 1, column=col).fill = fill(LEVELS[k][1])
        ws.cell(row=t0 + 1, column=col).font = Font(bold=True, color=LEVELS[k][2].lstrip("#"))
    for i, (tac, tname) in enumerate(attack.tactics, t0 + 2):
        in_tac = [r for r in rows if tac in r["tactics"]]
        vals = [tname] + [sum(1 for r in in_tac if r["level"] == k) for k in LEVEL_ORDER] + [len(in_tac)]
        for c, v in enumerate(vals, 1):
            ws.cell(row=i, column=c, value=v).border = BORDER
    for col, w in zip("ABCDEFG", (24, 40, 60, 20, 50, 14, 14)):
        ws.column_dimensions[col].width = w

    # ---- Technique Overlap (parents, with their sub-techniques nested underneath as collapsible rows)
    ws = wb.create_sheet("Technique Overlap")
    hdr = ["Overlap", "# Groups", "Type", "Technique ID", "Technique Name", "Tactics"] + [g["name"] for g in groups] + ["ATT&CK URL"]
    gcol = 7
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font, cell.border, cell.alignment = HDR, HFONT, BORDER, CENTER
    i = 1
    for r in rows:
        tactics = ", ".join(attack.tactic_name[t] for t in r["tactics"])
        for t, is_sub in [(r, False)] + [(sr, True) for sr in r["sub_rows"]]:
            i += 1
            lv = LEVELS[t["level"]]
            vals = [LEVEL_SHORT[t["level"]], f"{t['count']}/{n}", "Sub-technique" if is_sub else "Technique",
                    t["id"], t["name"], tactics] + ["✔" if g["id"] in t["groups"] else "" for g in groups] + [t["url"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=i, column=c, value=v)
                cell.border, cell.alignment = BORDER, WRAP
                if is_sub:
                    cell.font = Font(italic=True, color="555555")
            for c in (1, 2):
                cell = ws.cell(row=i, column=c)
                cell.fill, cell.alignment = fill(lv[1]), CENTER
                cell.font = Font(bold=not is_sub, italic=is_sub, color=lv[2].lstrip("#"))
            if is_sub:
                for c in (4, 5):
                    ws.cell(row=i, column=c).alignment = Alignment(wrap_text=True, vertical="top", indent=2)
                ws.row_dimensions[i].outlineLevel = 1
            else:
                for c in (4, 5):
                    ws.cell(row=i, column=c).font = Font(bold=True)
            for j, g in enumerate(groups):
                if g["id"] in t["groups"]:
                    cell = ws.cell(row=i, column=gcol + j)
                    cell.fill, cell.alignment = fill(lv[1]), CENTER
                    cell.font = Font(bold=True, color=lv[2].lstrip("#"))
            ws.cell(row=i, column=len(hdr)).hyperlink = t["url"]
    ws.sheet_properties.outlinePr.summaryBelow = False
    widths = [13, 9, 14, 13, 42, 30] + [16] * n + [48]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{i}"

    # ---- Tactic Matrix (Navigator-style)
    ws = wb.create_sheet("Tactic Matrix")
    used_tactics = [(t, nm) for t, nm in attack.tactics if any(t in r["tactics"] for r in rows)]
    for c, (tac, tname) in enumerate(used_tactics, 1):
        col_rows = [r for r in rows if tac in r["tactics"]]
        cell = ws.cell(row=1, column=c, value=f"{tname}\n({len(col_rows)})")
        cell.fill, cell.font, cell.border, cell.alignment = HDR, HFONT, BORDER, CENTER
        i = 1
        for r in col_rows:
            for t, is_sub in [(r, False)] + [(sr, True) for sr in r["sub_rows"]]:
                i += 1
                lv = LEVELS[t["level"]]
                label = f"↳ {t['id']}  {t['name']}  [{t['count']}/{n}]" if is_sub else f"{t['id']}  {t['name']}  [{t['count']}/{n}]"
                cell = ws.cell(row=i, column=c, value=label)
                cell.fill, cell.border = fill(lv[1]), BORDER
                cell.font = Font(color=lv[2].lstrip("#"), size=8 if is_sub else 9, bold=not is_sub, italic=is_sub)
                cell.alignment = Alignment(wrap_text=True, vertical="center", indent=1 if is_sub else 0)
        ws.column_dimensions[get_column_letter(c)].width = 24
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    wb.save(path)


def write_pdf(path, attack, groups, rows, per_group_counts, title, opts):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, TABLOID
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.utils import simpleSplit
    from reportlab.platypus import (BaseDocTemplate, Flowable, Frame, PageBreak, PageTemplate, Paragraph,
                                    Spacer, Table, TableStyle)

    PAGE = landscape(TABLOID)
    W, H = PAGE
    M = 0.45 * inch
    BG = colors.HexColor("#0E1117")
    PANEL = colors.HexColor("#171B24")
    PANEL2 = colors.HexColor("#1F2633")
    LINE = colors.HexColor("#2C3442")
    TXT = colors.HexColor("#E6E9EF")
    MUTED = colors.HexColor("#8B93A1")
    ACCENT = colors.HexColor("#5AA9E6")
    n = len(groups)
    gname = {g["id"]: g["name"] for g in groups}
    stamp = f"{dt.datetime.now():%Y-%m-%d %H:%M}"

    def on_page(c, doc):
        c.saveState()
        c.setFillColor(BG)
        c.rect(0, 0, W, H, stroke=0, fill=1)
        c.setFillColor(ACCENT)
        c.rect(0, H - 4, W, 4, stroke=0, fill=1)
        c.setFont("Helvetica", 7.5)
        c.setFillColor(MUTED)
        c.drawString(M, 0.25 * inch, f"MITRE ATT&CK Navigator for Dummies  |  {', '.join(gname.values())}"[:180])
        c.drawRightString(W - M, 0.25 * inch, f"Generated {stamp}  |  ATT&CK data {attack.data_modified}  |  Page {doc.page}")
        c.restoreState()

    doc = BaseDocTemplate(str(path), pagesize=PAGE, leftMargin=M, rightMargin=M, topMargin=M, bottomMargin=0.5 * inch,
                          title=title, author="MITRE ATT&CK Navigator for Dummies")
    frame = Frame(M, 0.5 * inch, W - 2 * M, H - M - 0.5 * inch, id="f", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="dark", frames=[frame], onPage=on_page)])

    sty = lambda name, **kw: ParagraphStyle(name, fontName=kw.pop("font", "Helvetica"), textColor=kw.pop("color", TXT), **kw)
    H1 = sty("h1", font="Helvetica-Bold", fontSize=26, leading=30)
    H2 = sty("h2", font="Helvetica-Bold", fontSize=15, leading=19, spaceBefore=6, spaceAfter=6, color=ACCENT)
    SUB = sty("sub", fontSize=10.5, leading=14, color=MUTED)
    CELL = sty("cell", fontSize=8.5, leading=10.5)
    CELLB = sty("cellb", font="Helvetica-Bold", fontSize=8.5, leading=10.5)
    HDRS = sty("hdr", font="Helvetica-Bold", fontSize=8.5, leading=10.5, color=colors.white)

    def dark_table(data, widths, extra=()):
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), PANEL2),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [PANEL, colors.HexColor("#141821")]),
            ("GRID", (0, 0), (-1, -1), 0.4, LINE),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            *extra,
        ]))
        return t

    counts = level_counts(rows)
    n_subs = sum(len(r["sub_rows"]) for r in rows)
    story = []

    # ---- Page 1: summary
    story += [Paragraph(title, H1), Spacer(1, 4),
              Paragraph(f"Technique overlap across {n} selected group{'s' if n != 1 else ''} &nbsp;|&nbsp; "
                        f"{len(rows)} parent T-codes" + (f" + {n_subs} sub-techniques" if opts["subs"] else " (sub-techniques rolled up)") +
                        f" &nbsp;|&nbsp; campaigns {'included' if opts['campaigns'] else 'excluded'} "
                        f"&nbsp;|&nbsp; software-derived {'included' if opts['software'] else 'excluded'}", SUB),
              Spacer(1, 14)]

    class Legend(Flowable):
        def __init__(self, width):
            super().__init__()
            self.width, self.height = width, 78

        def draw(self):
            c = self.canv
            gap = 14
            nb = len(LEVEL_ORDER)
            bw = (self.width - gap * (nb - 1)) / max(nb, 3)
            meaning = level_meanings(n)
            small = nb > 3
            for i, k in enumerate(LEVEL_ORDER):
                x = i * (bw + gap)
                c.setFillColor(PANEL)
                c.setStrokeColor(LINE)
                c.roundRect(x, 0, bw, self.height, 8, stroke=1, fill=1)
                c.setFillColor(colors.HexColor(LEVELS[k][1]))
                c.roundRect(x, 0, 10, self.height, 4, stroke=0, fill=1)
                c.rect(x + 5, 0, 5, self.height, stroke=0, fill=1)
                c.setFillColor(TXT if k == ALL else colors.HexColor(LEVELS[k][1]))  # dark red is unreadable on the dark panel
                big = 26 if small else 30
                c.setFont("Helvetica-Bold", big)
                num = str(counts[k][0])
                c.drawString(x + 22, self.height - 42, num)
                tx = x + 22 + c.stringWidth(num, "Helvetica-Bold", big) + 10
                c.setFillColor(TXT)
                c.setFont("Helvetica-Bold", 11 if small else 12)
                c.drawString(tx, self.height - 30, LEVEL_SHORT[k])
                c.setFillColor(MUTED)
                c.setFont("Helvetica", 8.5 if small else 9.5)
                c.drawString(tx, self.height - 44, (f"tech  +  {counts[k][1]} sub" if small else
                                                    f"techniques  +  {counts[k][1]} sub-techniques") if opts["subs"] else "techniques")
                c.setFillColor(TXT)
                c.setFont("Helvetica", 9.5 if small else 10)
                c.drawString(x + 22, 16, meaning[k])

    story += [Legend(W - 2 * M), Spacer(1, 16)]

    left_w = (W - 2 * M) * 0.48
    right_w = (W - 2 * M) - left_w - 18

    gdata = [[Paragraph(h, HDRS) for h in ("ID", "Group", "Aliases", "T-codes")]]
    for g in groups:
        al = ", ".join(g["aliases"])
        gdata.append([Paragraph(g["id"], CELLB), Paragraph(g["name"], CELLB),
                      Paragraph(al[:220] + ("..." if len(al) > 220 else ""), CELL), Paragraph(str(per_group_counts[g["id"]]), CELL)])
    gtable = dark_table(gdata, [0.6 * inch, 1.45 * inch, left_w - 2.75 * inch, 0.7 * inch])

    tdata = [[Paragraph("Tactic", HDRS)] + [Paragraph(LEVEL_SHORT[k], sty("h" + k, font="Helvetica-Bold", fontSize=8.5, leading=10.5,
                                                                        color=colors.HexColor(LEVELS[k][2]))) for k in LEVEL_ORDER]
             + [Paragraph("Total", HDRS)]]
    for tac, tname in attack.tactics:
        in_tac = [r for r in rows if tac in r["tactics"]]
        tdata.append([Paragraph(tname, CELL)] + [Paragraph(str(sum(1 for r in in_tac if r["level"] == k)), CELL) for k in LEVEL_ORDER]
                     + [Paragraph(str(len(in_tac)), CELLB)])
    cw = (right_w - 1.9 * inch) / (len(LEVEL_ORDER) + 1)
    ttable = dark_table(tdata, [1.9 * inch] + [cw] * (len(LEVEL_ORDER) + 1),
                        extra=[("BACKGROUND", (1 + i, 0), (1 + i, 0), colors.HexColor(LEVELS[k][1])) for i, k in enumerate(LEVEL_ORDER)])

    side = Table([[[Paragraph("Selected Groups", H2), gtable], "", [Paragraph("Coverage by Tactic", H2), ttable]]],
                 colWidths=[left_w, 18, right_w])
    side.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story += [side, PageBreak()]

    # ---- Matrix pages
    used_tactics = [(t, nm) for t, nm in attack.tactics if any(t in r["tactics"] for r in rows)]
    cols = {t: [r for r in rows if t in r["tactics"]] for t, _ in used_tactics}
    CELL_H, SUB_H, GAP_Y = 25, 12, 2
    avail_h = H - M - 0.5 * inch
    capacity = avail_h - 2 - 30 - 30 - 6   # page height minus title row and tactic header

    def paginate(col_rows):
        """Split one tactic column into pages, keeping a parent together with its sub-techniques when it fits."""
        pages, cur, used = [], [], 0
        for r in col_rows:
            block = [(r, False)] + [(sr, True) for sr in r["sub_rows"]]
            block_h = sum((SUB_H if sub else CELL_H) + GAP_Y for _, sub in block)
            if cur and used + block_h > capacity and block_h <= capacity:
                pages.append(cur)
                cur, used = [], 0
            for item in block:
                h = (SUB_H if item[1] else CELL_H) + GAP_Y
                if used + h > capacity:
                    pages.append(cur)
                    cur, used = [], 0
                cur.append(item)
                used += h
        if cur:
            pages.append(cur)
        return pages

    col_pages = {t: paginate(cols[t]) for t, _ in used_tactics}
    n_pages = max([len(v) for v in col_pages.values()] + [1])

    def fit(c, text, font, size, width):
        if c.stringWidth(text, font, size) <= width:
            return text
        while text and c.stringWidth(text + "...", font, size) > width:
            text = text[:-1]
        return text + "..."

    class MatrixPage(Flowable):
        """Navigator-style matrix: one column per tactic, parent cells with their sub-technique cells indented below."""
        CELL_FONT = 6.3

        def __init__(self, width, height, page_idx):
            super().__init__()
            self.width, self.height, self.page_idx = width, height, page_idx

        def draw(self):
            c = self.canv
            ncol = len(used_tactics)
            gap = 3
            cw = (self.width - gap * (ncol - 1)) / ncol
            top = self.height
            label = "Technique Matrix" + (f"  ({self.page_idx + 1}/{n_pages})" if n_pages > 1 else "")
            c.setFillColor(TXT)
            c.setFont("Helvetica-Bold", 15)
            c.drawString(0, top - 15, label)
            c.setFillColor(MUTED)
            c.setFont("Helvetica", 8.5)
            c.drawString(c.stringWidth(label, "Helvetica-Bold", 15) + 14, top - 13,
                         "indented rows = sub-techniques, each colored by its own overlap" if opts["subs"]
                         else "parent T-codes only (sub-technique use rolled up into the parent)")
            # inline legend
            lx = self.width
            for k in reversed(LEVEL_ORDER):
                txt = f"{LEVEL_SHORT[k]} ({counts[k][0]} + {counts[k][1]} sub)" if opts["subs"] else f"{LEVEL_SHORT[k]} ({counts[k][0]})"
                c.setFont("Helvetica", 8.5)
                lx -= c.stringWidth(txt, "Helvetica", 8.5)
                c.setFillColor(TXT)
                c.drawString(lx, top - 13, txt)
                lx -= 14
                c.setFillColor(colors.HexColor(LEVELS[k][1]))
                c.roundRect(lx, top - 14, 10, 10, 2, stroke=0, fill=1)
                lx -= 16
            y0 = top - 30
            hh = 30
            for i, (tac, tname) in enumerate(used_tactics):
                x = i * (cw + gap)
                c.setFillColor(PANEL2)
                c.roundRect(x, y0 - hh, cw, hh, 3, stroke=0, fill=1)
                c.setFillColor(colors.white)
                lines = simpleSplit(tname, "Helvetica-Bold", 7.5, cw - 6)
                ty = y0 - 11 if len(lines) > 1 else y0 - 14
                for ln in lines[:2]:
                    c.setFont("Helvetica-Bold", 7.5)
                    c.drawCentredString(x + cw / 2, ty, ln)
                    ty -= 8.5
                c.setFillColor(MUTED)
                c.setFont("Helvetica", 6.5)
                c.drawCentredString(x + cw / 2, y0 - hh + 3.5, f"{len(cols[tac])} techniques")
                y = y0 - hh - 4
                pages = col_pages[tac]
                for t, is_sub in (pages[self.page_idx] if self.page_idx < len(pages) else []):
                    lv = LEVELS[t["level"]]
                    if is_sub:
                        sx, sw = x + 6, cw - 6
                        c.setFillColor(colors.HexColor(lv[1]))
                        c.roundRect(sx, y - SUB_H, sw, SUB_H, 1.5, stroke=0, fill=1)
                        c.setFillColor(colors.HexColor(lv[2]))
                        sid = "." + t["id"].split(".")[1]
                        cnt = f"{t['count']}/{n}"
                        c.setFont("Helvetica-Bold", 5.6)
                        c.drawString(sx + 2.5, y - 8.3, sid)
                        c.setFont("Helvetica", 5.3)
                        c.drawRightString(sx + sw - 2.5, y - 8.3, cnt)
                        nx = sx + 2.5 + c.stringWidth(sid, "Helvetica-Bold", 5.6) + 3
                        room = sx + sw - 2.5 - c.stringWidth(cnt, "Helvetica", 5.3) - 3 - nx
                        c.setFont("Helvetica", 5.6)
                        c.drawString(nx, y - 8.3, fit(c, t["name"], "Helvetica", 5.6, room))
                        y -= SUB_H + GAP_Y
                        continue
                    c.setFillColor(colors.HexColor(lv[1]))
                    c.roundRect(x, y - CELL_H, cw, CELL_H, 2, stroke=0, fill=1)
                    c.setFillColor(colors.HexColor(lv[2]))
                    c.setFont("Helvetica-Bold", 6.8)
                    c.drawString(x + 3, y - 8, t["id"])
                    c.setFont("Helvetica", 6)
                    c.drawRightString(x + cw - 3, y - 8, f"{t['count']}/{n}")
                    nl = simpleSplit(t["name"], "Helvetica", self.CELL_FONT, cw - 6)
                    if len(nl) > 2:
                        nl = [nl[0], fit(c, " ".join(nl[1:]), "Helvetica", self.CELL_FONT, cw - 6)]
                    c.setFont("Helvetica", self.CELL_FONT)
                    ly = y - 15.5
                    for ln in nl:
                        c.drawString(x + 3, ly, ln)
                        ly -= 7
                    y -= CELL_H + GAP_Y

    for p in range(n_pages):
        story.append(MatrixPage(W - 2 * M, avail_h - 2, p))
        story.append(PageBreak())

    # ---- Detail table (each parent followed by its sub-techniques)
    story.append(Paragraph("Technique &amp; Sub-Technique Detail" if opts["subs"] else "Technique Detail", H2))
    widths = [0.95 * inch, 0.6 * inch, 0.85 * inch, 0, 2.2 * inch]
    free = (W - 2 * M) - sum(widths)
    gcol_w = min(1.1 * inch, (free - 2.4 * inch) / max(n, 1))
    widths[3] = free - gcol_w * n
    hdr = ["Overlap", "Groups", "ID", "Technique / Sub-technique", "Tactics"] + [gname[g["id"]] for g in groups]
    data = [[Paragraph(h, HDRS) for h in hdr]]
    extra = []
    SUBID = sty("subid", font="Helvetica-Bold", fontSize=7.5, leading=9.5, color=MUTED, leftIndent=8)
    SUBNM = sty("subnm", fontSize=7.8, leading=9.5, color=colors.HexColor("#C3C8D2"), leftIndent=8)
    i = 0
    for r in rows:
        for t, is_sub in [(r, False)] + [(sr, True) for sr in r["sub_rows"]]:
            i += 1
            lv = LEVELS[t["level"]]
            fc = colors.HexColor(lv[2])
            fs = 7 if is_sub else 8
            row = [Paragraph(LEVEL_SHORT[t["level"]], sty(f"lv{i}", font="Helvetica-Oblique" if is_sub else "Helvetica-Bold",
                                                          fontSize=fs, leading=fs + 2, color=fc, leftIndent=6 if is_sub else 0)),
                   Paragraph(f"{t['count']}/{n}", sty(f"ct{i}", font="Helvetica-Bold", fontSize=fs + 0.5, leading=fs + 2.5, color=fc)),
                   Paragraph(t["id"], SUBID if is_sub else CELLB),
                   Paragraph(t["name"], SUBNM if is_sub else CELLB),
                   Paragraph("" if is_sub else ", ".join(attack.tactic_name[x] for x in r["tactics"]),
                             sty(f"tc{i}", fontSize=7.5, leading=9.5, color=MUTED))]
            for j, g in enumerate(groups):
                if g["id"] in t["groups"]:
                    row.append(Paragraph("✔", sty(f"ck{i}{j}", font="ZapfDingbats", fontSize=8 if is_sub else 9,
                                                       leading=10, alignment=1, color=fc)))
                    extra.append(("BACKGROUND", (5 + j, i), (5 + j, i), colors.HexColor(lv[1])))
                else:
                    row.append("")
            data.append(row)
            extra.append(("BACKGROUND", (0, i), (1, i), colors.HexColor(lv[1])))
            if is_sub:
                extra += [("TOPPADDING", (0, i), (-1, i), 1.5), ("BOTTOMPADDING", (0, i), (-1, i), 1.5),
                          ("BACKGROUND", (2, i), (4, i), colors.HexColor("#12151C"))]
            else:
                extra.append(("LINEABOVE", (0, i), (-1, i), 1.1, colors.HexColor("#3A4456")))
    story.append(dark_table(data, widths + [gcol_w] * n, extra=extra))

    doc.build(story)


_LEGEND_FONT = None
# Libre Franklin: open-source cut of Franklin Gothic (OFL). The variable font renders the PNG legend (at its Bold
# instance); the static Bold copy is what reportlab can read to measure text for the pure-Python SVG legend.
BUNDLED_FONT = HERE / "fonts" / "LibreFranklin-Bold.ttf"
BUNDLED_VARIABLE_FONT = HERE / "fonts" / "LibreFranklin-VariableFont.ttf"
SVG_FONT_FAMILY = "'Franklin Gothic Medium', 'Franklin Gothic', 'Libre Franklin', Arial, sans-serif"


def legend_font_source():
    """Font for the PNG legend: Franklin Gothic (bold weight) from the system fonts, else the bundled Libre Franklin.

    Returns (path, variation name or None, display name). Only called when Pillow is available.
    """
    global _LEGEND_FONT
    if _LEGEND_FONT:
        return _LEGEND_FONT
    from PIL import ImageFont
    dirs = [Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts",
            Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "Windows" / "Fonts",
            Path("/Library/Fonts"), Path.home() / "Library" / "Fonts", Path("/usr/share/fonts"), Path.home() / ".fonts"]
    rank = ["bold", "demi", "heavy", "medium"]   # closest to "Franklin Gothic Bold" first
    best = None
    for d in dirs:
        if not d.is_dir():
            continue
        for f in d.rglob("*"):
            if f.suffix.lower() not in (".ttf", ".otf") or "fra" not in f.name.lower():
                continue
            try:
                fam, style = ImageFont.truetype(str(f), 10).getname()
            except OSError:
                continue
            full = f"{fam} {style}".lower()
            if "franklin gothic" not in full or "italic" in full or "oblique" in full or "cond" in full:
                continue
            r = next((i for i, w in enumerate(rank) if w in full), None)
            if r is not None and (best is None or r < best[0]):
                best = (r, str(f), None, f"{fam} {style}".replace(" Regular", ""))
    _LEGEND_FONT = best[1:] if best else (str(BUNDLED_VARIABLE_FONT), "Bold",
                                          "Libre Franklin Bold (Franklin Gothic not installed)")
    return _LEGEND_FONT


def legend_label(lo, hi, n):
    if lo == hi == n and n > 1:
        return "Both groups" if n == 2 else f"All {n} groups"
    if lo == hi:
        return "1 group" if lo == 1 else f"{lo} of {n} groups"
    return f"{lo}\u2013{hi} of {n} groups"


LEGEND_THEMES = {   # HEAD, CELL, TXT, LINE, EDGE
    # for the dark Navigator: dark gray cells, light text and a light outline so it stands off the page
    "dark": ("#3F3F3F", "#2B2B2B", "#FFFFFF", "#7F7F7F", "#BFBFBF"),
    # for the light Navigator: light gray cells, black text, dark outline against the white page
    "light": ("#BFBFBF", "#D9D9D9", "#000000", "#7F7F7F", "#404040"),
}
HEAD_PT, CELL_PT = 68, 54


def legend_layout(groups, measure):
    """Shared geometry for the PNG and SVG legends. measure(text, size) -> width in px."""
    n = len(groups)
    rows = [(lv, legend_label(lo, hi, n)) for lv, lo, hi in BANDS]
    sw_w, row_h, head_h, tpad, edge = 170, 112, 124, 44, 6
    lab_w = int(max(measure(t, CELL_PT) for _, t in rows)) + 2 * tpad
    width = max(sw_w + lab_w, int(measure("LEGEND", HEAD_PT)) + 2 * tpad) + 2 * edge
    height = head_h + row_h * len(rows) + 2 * edge
    return dict(rows=rows, sw_w=sw_w, row_h=row_h, head_h=head_h, tpad=tpad, edge=edge, width=width, height=height,
                x0=edge, x1=edge + sw_w, x2=width - edge)


def write_legend_png(path, groups, theme):
    """Slide legend as PNG (needs Pillow). One table: a merged LEGEND header, then a color cell and label per level."""
    from PIL import Image, ImageDraw, ImageFont
    src, var, _ = legend_font_source()
    fonts = {sz: ImageFont.truetype(src, sz) for sz in (HEAD_PT, CELL_PT)}
    if var:
        for f in fonts.values():
            f.set_variation_by_name(var)
    L = legend_layout(groups, lambda t, sz: fonts[sz].getlength(t))
    HEAD, CELL, TXT, LINE, EDGE = LEGEND_THEMES[theme]
    x0, x1, x2, edge, width, height = L["x0"], L["x1"], L["x2"], L["edge"], L["width"], L["height"]

    img = Image.new("RGB", (width, height), EDGE)
    d = ImageDraw.Draw(img)
    d.rectangle([x0, edge, x2 - 1, edge + L["head_h"]], fill=HEAD)
    d.text((width / 2, edge + L["head_h"] / 2), "LEGEND", font=fonts[HEAD_PT], fill=TXT, anchor="mm")
    y = edge + L["head_h"]
    for lv, text in L["rows"]:
        d.line([x0, y, x2, y], fill=LINE, width=3)
        d.rectangle([x0, y + 2, x1 - 2, y + L["row_h"]], fill=LEVELS[lv][1])
        d.rectangle([x1 + 2, y + 2, x2 - 1, y + L["row_h"]], fill=CELL)
        d.text((x1 + L["tpad"], y + L["row_h"] / 2 + 1), text, font=fonts[CELL_PT], fill=TXT, anchor="lm")
        y += L["row_h"]
    d.line([x1, edge + L["head_h"], x1, height - edge], fill=LINE, width=3)
    img.save(path, dpi=(300, 300))


def write_legend_svg(path, groups, theme):
    """Slide legend as SVG, pure Python (no Pillow). Same table as the PNG; vector, so it stays sharp at any size.

    Text uses the workstation's Franklin Gothic when PowerPoint renders it. Widths are measured with the bundled
    Libre Franklin Bold plus 8% headroom, since installed Franklin Gothic cuts run slightly wider.
    """
    from xml.sax.saxutils import escape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if "LibreFranklin-Bold" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("LibreFranklin-Bold", str(BUNDLED_FONT)))
    face = pdfmetrics.getFont("LibreFranklin-Bold").face
    cap = getattr(face, "capHeight", 700) / 1000   # cap height as a fraction of the font size, to center text
    L = legend_layout(groups, lambda t, sz: pdfmetrics.stringWidth(t, "LibreFranklin-Bold", sz) * 1.08)
    HEAD, CELL, TXT, LINE, EDGE = LEGEND_THEMES[theme]
    x0, x1, x2, edge, width, height = L["x0"], L["x1"], L["x2"], L["edge"], L["width"], L["height"]
    base = lambda cy, sz: round(cy + cap * sz / 2, 1)   # baseline that vertically centers capital letters on cy
    font = f'font-family="{SVG_FONT_FAMILY}" font-weight="bold"'

    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
           f'<rect x="0" y="0" width="{width}" height="{height}" fill="{EDGE}"/>',
           f'<rect x="{x0}" y="{edge}" width="{x2 - x0}" height="{L["head_h"]}" fill="{HEAD}"/>',
           f'<text x="{width / 2}" y="{base(edge + L["head_h"] / 2, HEAD_PT)}" {font} font-size="{HEAD_PT}" '
           f'fill="{TXT}" text-anchor="middle">LEGEND</text>']
    y = edge + L["head_h"]
    for lv, text in L["rows"]:
        out += [f'<rect x="{x0}" y="{y}" width="{x1 - x0}" height="{L["row_h"]}" fill="{LEVELS[lv][1]}"/>',
                f'<rect x="{x1}" y="{y}" width="{x2 - x1}" height="{L["row_h"]}" fill="{CELL}"/>',
                f'<line x1="{x0}" y1="{y}" x2="{x2}" y2="{y}" stroke="{LINE}" stroke-width="3"/>',
                f'<text x="{x1 + L["tpad"]}" y="{base(y + L["row_h"] / 2, CELL_PT)}" {font} font-size="{CELL_PT}" '
                f'fill="{TXT}">{escape(text)}</text>']
        y += L["row_h"]
    out += [f'<line x1="{x1}" y1="{edge + L["head_h"]}" x2="{x1}" y2="{height - edge}" stroke="{LINE}" stroke-width="3"/>',
            "</svg>"]
    Path(path).write_text("\n".join(out) + "\n", encoding="utf-8")


# --------------------------------------------------------------------------- CLI

def print_groups(groups):
    for g in groups:
        al = ", ".join(g["aliases"])
        print(f"  {g['id']:<6} {g['name']:<32} {al[:80]}")


def resolve(attack, queries):
    chosen, problems = [], []
    for q in queries:
        g, sugg = attack.find_group(q)
        if g:
            if g["id"] not in {c["id"] for c in chosen}:
                chosen.append(g)
        else:
            problems.append((q, sugg))
    return chosen, problems


def ask_domain():
    while True:
        try:
            line = input("Matrix? Press Enter for Enterprise, or type mobile / ics: ").strip()
        except EOFError:
            return None
        if not line:
            return "enterprise"
        if line.lower() in ("q", "quit", "exit"):
            return None
        d = parse_domain(line)
        if d:
            return d
        print("  Please type enterprise, mobile or ics (or just press Enter for Enterprise).")


def interactive(attack, domain):
    print(f"\n[{DOMAIN_LABEL[domain]} matrix: {len(attack.groups)} groups]")
    print("Enter APT groups separated by commas. Use the name, an alias, or a G-ID.")
    print("  e.g.  APT28, APT29, Lazarus Group      or     G0007, G0016, G0032")
    print("Commands:  list   = show all groups")
    print("           find X = search groups by name/alias (e.g. find bear)")
    print("           quit   = exit\n")
    while True:
        try:
            line = input("Groups> ").strip()
        except EOFError:
            return None
        if not line:
            continue
        low = line.lower()
        if low in ("q", "quit", "exit"):
            return None
        if low == "list":
            print_groups(sorted(attack.groups.values(), key=lambda g: g["name"].lower()))
            continue
        if low.startswith(("find ", "search ")):
            hits = attack.search(line.split(" ", 1)[1])
            print_groups(hits) if hits else print("  No matches.")
            continue
        queries = [q for q in (s.strip() for s in line.split(",")) if q]
        chosen, problems = resolve(attack, queries)
        for q, sugg in problems:
            print(f"\n  Could not uniquely match '{q}'.")
            if sugg:
                print("  Did you mean one of these? (use the G-ID to be exact)")
                print_groups(sugg[:15])
            else:
                print("  No similar groups found. Try 'find <word>'.")
        if problems:
            print()
            continue
        print("\nSelected:")
        print_groups(chosen)
        ok = input("\nGenerate layer for these groups? [Y/n] ").strip().lower()
        if ok in ("", "y", "yes"):
            return chosen


def main():
    ap = argparse.ArgumentParser(description="Generate color-coded ATT&CK Navigator overlap layers (JSON + Excel + dark PDF) for APT groups.")
    ap.add_argument("groups", nargs="*", metavar="[mobile|ics] GROUP",
                    help="optional matrix (mobile or ics; Enterprise if left out) followed by group names, "
                         "aliases or G-IDs (omit everything for interactive mode)")
    ap.add_argument("--list", action="store_true", help="list all groups and exit")
    ap.add_argument("--search", metavar="TEXT", help="search groups by name/alias and exit")
    ap.add_argument("--no-campaigns", action="store_true", help="exclude techniques from campaigns attributed to the groups")
    sw = ap.add_mutually_exclusive_group()
    sw.add_argument("--include-software", dest="software", action="store_true", default=None,
                    help="also include techniques of malware/tools the groups use (default for mobile and ics)")
    sw.add_argument("--no-software", dest="software", action="store_false",
                    help="only count techniques linked directly to the group (default for enterprise)")
    ap.add_argument("--mode", choices=["both", "parents", "subs"], default="both",
                    help="which outputs to build: parent-only, parent+sub-techniques, or both (default)")
    ap.add_argument("--domain", choices=list(DOMAINS), help=argparse.SUPPRESS)  # older spelling of the matrix word
    ap.add_argument("--update", action="store_true", help="re-download the latest ATT&CK data first")
    ap.add_argument("--offline", action="store_true",
                    help="never use the network: run only from the ATT&CK data bundled in data/")
    ap.add_argument("--title", help="custom report title")
    ap.add_argument("--out", help="output folder (default: ./output/<timestamp>_<groups>)")
    args = ap.parse_args()
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")  # Windows consoles choke on some alias characters
        except AttributeError:
            pass

    # Optional first word picks the matrix: "mobile" or "ics" (Enterprise otherwise).
    domain = args.domain
    if args.groups and parse_domain(args.groups[0]):
        domain = parse_domain(args.groups.pop(0))
    is_interactive = not args.groups and not (args.list or args.search)
    if is_interactive:
        print("\n" + "=" * 70)
        print("  MITRE ATT&CK NAVIGATOR FOR DUMMIES")
        print("=" * 70)
    if not domain:
        domain = ask_domain() if is_interactive else "enterprise"
        if not domain:
            return 0
    label = DOMAIN_LABEL[domain]
    if args.software is None:
        # Mobile/ICS groups are mostly linked to techniques through their malware, so count software there by default.
        args.software = domain != "enterprise"

    if args.offline and args.update:
        ap.error("--offline and --update cannot be used together")
    bundle, domain_name, data_path = load_bundle(domain, args.update, args.offline)
    attack = Attack(bundle, domain_name)
    print(f"ATT&CK {label} data: {data_path.relative_to(HERE).as_posix()} (MITRE data as of {attack.data_modified})"
          + ("  [offline]" if args.offline else ""))
    del bundle

    if args.list:
        print_groups(sorted(attack.groups.values(), key=lambda g: g["name"].lower()))
        return 0
    if args.search:
        hits = attack.search(args.search)
        print_groups(hits) if hits else print("No matches.")
        return 0

    if is_interactive:
        groups = interactive(attack, domain)
        if not groups:
            return 0
    else:
        groups, problems = resolve(attack, args.groups)
        if problems:
            for q, sugg in problems:
                print(f"Could not uniquely match '{q}' in the {label} matrix.")
                if sugg:
                    print("Did you mean:")
                    print_groups(sugg[:15])
            print(f"\nNot every group is in every matrix. See the {label} groups with:  "
                  f"python attack_overlap.py {'' if domain == 'enterprise' else domain + ' '}--list")
            return 2

    opts = {"campaigns": not args.no_campaigns, "software": args.software}
    rows, per_group_counts = analyze(attack, groups, **opts)
    if not rows:
        print(f"The selected groups have no recorded techniques in the {label} matrix.")
        return 1
    for g in groups:
        if not per_group_counts.get(g["id"]):
            print(f"Note: {g['name']} ({g['id']}) has no recorded techniques in the {label} matrix"
                  + ("" if opts["software"] else " (try --include-software)") + ".")

    title = args.title or (f"APT Overlap ({label}): " + " vs ".join(g["name"] for g in groups))
    base = safe_name(label + "_" + "_".join(g["name"] for g in groups))[:80]
    out = Path(args.out) if args.out else OUT_DIR / f"{dt.datetime.now():%Y%m%d_%H%M%S}_{base}"

    counts = level_counts(rows)
    n_subs = sum(len(r["sub_rows"]) for r in rows)
    print("\nColor scale:  " + "  |  ".join(f"{LEVEL_SHORT[k]} = {m.replace('Used by ', '')}" for k, m in level_meanings(len(groups)).items()))
    print(f"{len(rows)} parent T-codes:  " + "  |  ".join(f"{LEVEL_SHORT[k]} {counts[k][0]}" for k in LEVEL_ORDER))
    print(f"{n_subs} sub-techniques:  " + "  |  ".join(f"{LEVEL_SHORT[k]} {counts[k][1]}" for k in LEVEL_ORDER))

    # Each run produces two sets: parent T-codes only, and parents with their sub-techniques.
    variants = {"parents": ("1_Parent_Only", "parent-only", " (Parent T-codes)",
                            [dict(r, sub_rows=[]) for r in rows]),
                "subs": ("2_Parent_and_Sub", "parent-and-sub", " (Parent + Sub-techniques)", rows)}
    chosen = list(variants) if args.mode == "both" else [args.mode]
    out.mkdir(parents=True, exist_ok=True)
    legends = []
    for theme in ("light", "dark"):
        legends.append(out / f"{base}_legend_for-{theme}-navigator.svg")
        write_legend_svg(legends[-1], groups, theme)
        if HAVE_PILLOW:
            legends.append(out / f"{base}_legend_for-{theme}-navigator.png")
            write_legend_png(legends[-1], groups, theme)
    for key in chosen:
        folder, suffix, title_suffix, vrows = variants[key]
        vout = out / folder
        vout.mkdir(parents=True, exist_ok=True)
        vopts = dict(opts, subs=(key == "subs"))
        vtitle = title + title_suffix
        layer_p = vout / f"{base}_{suffix}_layer.json"
        xlsx_p = vout / f"{base}_{suffix}.xlsx"
        pdf_p = vout / f"{base}_{suffix}.pdf"
        write_layer(layer_p, attack, groups, vrows, vtitle, vopts)
        write_excel(xlsx_p, attack, groups, vrows, per_group_counts, vtitle, vopts)
        write_pdf(pdf_p, attack, groups, vrows, per_group_counts, vtitle, vopts)
        print(f"\n[{folder}]")
        print(f"  Navigator layer : {layer_p}")
        print(f"  Excel workbook  : {xlsx_p}")
        print(f"  PDF report      : {pdf_p}")
    print("\n[Slide legend]  " + (f"PNG font: {legend_font_source()[2]}" if HAVE_PILLOW
                                   else "SVG only (Pillow not installed, so no PNG copies)"))
    for lp in legends:
        print(f"  {lp.suffix[1:].upper():<15} : {lp}")
    print("\nTip: open https://mitre-attack.github.io/attack-navigator/ -> Open Existing Layer -> Upload from local -> pick the _layer.json")

    if is_interactive and sys.platform == "win32":
        os.startfile(out)
        input("\nPress Enter to close...")
    return 0


if __name__ == "__main__":
    sys.exit(main())
